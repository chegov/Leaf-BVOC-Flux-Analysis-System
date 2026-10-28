"""Reproducible VOC hourly analysis. All timestamps are local, naive datetimes."""
from __future__ import annotations

import argparse
import hashlib
import json
import shutil
import unicodedata
from pathlib import Path

import numpy as np
import pandas as pd
from scipy import stats
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import matplotlib.dates as mdates
from flux import FLUX_UNIT, area_from_dataset_name, flux_settings, calculate_flux_rows
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

ROOT = Path(__file__).resolve().parent
INPUT_ROOT = ROOT.parent / 'Original Data'
CATEGORIES = ['iso', 'MTs', 'SQTs', 'OTs', 'OMTs', 'OVOCs']
PLOT_LABELS = {
    'temperature': 'Temperature', 'humidity': 'Relative humidity', 'PAR': 'PAR',
    'iso': 'Isoprene (iso)', 'MTs': 'Monoterpene hydrocarbons (MTs)',
    'SQTs': 'Sesquiterpene hydrocarbons (SQTs)', 'OTs': 'Other hydrocarbons (OTs)',
    'OMTs': 'Oxygenated monoterpenes (OMTs)', 'OVOCs': 'Other oxygenated VOCs (OVOCs)'}


def find_source_pair(folder):
    """Find a matching concentration workbook and one *_env.xlsx in a dataset folder."""
    folder = Path(folder)
    if not folder.is_dir():
        raise ValueError(f'原始数据文件夹不存在: {folder}')
    files = sorted(p for p in folder.iterdir()
                   if p.is_file() and p.suffix.lower() == '.xlsx' and not p.name.startswith('~$'))
    env_files = [p for p in files if p.stem.lower().endswith('_env')]
    voc_files = [p for p in files if p not in env_files]
    if len(env_files) != 1 or len(voc_files) != 1:
        raise ValueError(f'{folder.name}: 需要一个浓度xlsx和一个 *_env.xlsx；'
                         f'发现浓度文件{len(voc_files)}个、环境文件{len(env_files)}个，请检查目录或用 --voc/--env 指定。')
    return voc_files[0], env_files[0]


def summarize(frame, keys, columns):
    """Tidy sample mean and SE, with per-variable valid n; n<2 gives NaN SE."""
    rows = []
    for key, group in frame.groupby(keys, sort=True):
        key = key if isinstance(key, tuple) else (key,)
        for col in columns:
            x = group[col].dropna()
            rows.append(dict(zip(keys, key)) | {
                'variable': col, 'mean': x.mean(),
                'se': x.std(ddof=1) / np.sqrt(len(x)) if len(x) > 1 else np.nan,
                'n': len(x)})
    return pd.DataFrame(rows, columns=keys + ['variable','mean','se','n'])


def category_totals(data, mapping):
    out = data[['timestamp', 'hour']].copy()
    for cat in CATEGORIES:
        cols = mapping.loc[mapping.category.eq(cat), 'species'].tolist()
        # A missing constituent makes the total unavailable, never a silent zero.
        out[cat] = data[cols].sum(axis=1, min_count=len(cols)) if cols else np.nan
    return out


def match_environment(voc, env, mode, minutes=5):
    """Union per-hour windows; an environmental record is counted once per hour."""
    details, selected = [], []
    delta = pd.Timedelta(minutes=minutes)
    for hour, group in voc.groupby('hour', sort=True):
        mask = pd.Series(False, index=env.index)
        if mode == 'five_minute_bins':
            starts = group.timestamp.dt.floor(f'{minutes}min').drop_duplicates().sort_values()
            windows = [(t, t + delta) for t in starts]
        elif mode == 'plus_minus':
            windows = [(t - delta, t + delta) for t in group.timestamp]
        else:
            raise ValueError(f'Unknown window_mode: {mode}')
        for start, end in windows:
            current = (env.timestamp.ge(start) &
                       (env.timestamp.lt(end) if mode == 'five_minute_bins' else env.timestamp.le(end)))
            mask |= current
            details.append({'hour': hour, 'window_start': start, 'window_end': end,
                            'end_inclusive': mode == 'plus_minus', 'matched_n': int(current.sum())})
        part = env.loc[mask].copy()
        part.insert(0, 'hour', hour)
        selected.append(part)
    return pd.DataFrame(details), pd.concat(selected, ignore_index=True)


def bh_adjust(pvalues):
    p = np.asarray(pvalues, dtype=float)
    q = np.full(len(p), np.nan)
    good = np.flatnonzero(np.isfinite(p))
    order = good[np.argsort(p[good])]
    if len(order):
        q[order] = np.minimum(1, np.minimum.accumulate(
            (p[order] * len(order) / np.arange(1, len(order)+1))[::-1])[::-1])
    return q


def correlations(voc_stats, env_stats, level):
    v = voc_stats.pivot(index='hour', columns='variable', values='mean')
    e = env_stats.pivot(index='hour', columns='variable', values='mean')
    rows = []
    for species in v.columns:
        for parameter in e.columns:
            pair = pd.concat([v[species], e[parameter]], axis=1).dropna()
            for method, func in [('Pearson', stats.pearsonr), ('Spearman', stats.spearmanr)]:
                n = len(pair)
                status = 'ok'
                r = p = np.nan
                if n < 3:
                    status = 'insufficient_pairs'
                elif pair.iloc[:, 0].nunique() < 2 or pair.iloc[:, 1].nunique() < 2:
                    status = 'constant_series'
                else:
                    r, p = func(pair.iloc[:, 0], pair.iloc[:, 1])
                rows.append(dict(level=level, variable=species, environment=parameter,
                                 method=method, n=n, r=r, p=p, status=status))
    out = pd.DataFrame(rows)
    out['q_BH'] = np.nan
    for _, indices in out.groupby('method').groups.items():
        out.loc[indices, 'q_BH'] = bh_adjust(out.loc[indices, 'p'])
    return out


def read_inputs(voc_path, env_path, cfg, mapping_path):
    raw = pd.read_excel(voc_path, sheet_name=cfg['voc_sheet'])
    raw.columns = [str(c).strip() for c in raw.columns]
    mapping = pd.read_csv(mapping_path, keep_default_na=False)
    if mapping.species.duplicated().any():
        raise ValueError('分类表存在重复物种')
    if not set(mapping.category).issubset(set(CATEGORIES + ['excluded', 'review'])):
        raise ValueError('分类表 category 无效')
    named = [c for c in raw.columns if c not in [cfg['date_column'], cfg['time_column']]
             and c and not c.startswith('Unnamed:') and not c.startswith('.')]
    unknown = sorted(set(named) - set(mapping.species))
    if unknown:
        raise ValueError(f'新增物种尚未映射，请更新 species_mapping.csv: {unknown}')
    mapping = mapping[mapping.species.isin(raw.columns)].copy()
    data = pd.DataFrame()
    dates = pd.to_datetime(raw[cfg['date_column']], errors='raise').dt.normalize()
    times = raw[cfg['time_column']].map(lambda x: x.strftime('%H:%M:%S') if hasattr(x, 'strftime') else str(x).split()[-1])
    data['timestamp'] = dates + pd.to_timedelta(times)
    data['hour'] = data.timestamp.dt.floor('h')
    numeric = raw[mapping.species.tolist()].apply(pd.to_numeric, errors='raise')
    if np.isinf(numeric.to_numpy(dtype=float)).any():
        raise ValueError('浓度含无穷值，请检查原始数据')
    data = pd.concat([data, numeric], axis=1)
    if data.timestamp.isna().any() or data.timestamp.duplicated().any():
        raise ValueError('VOCs 时间缺失或重复，请先检查原始数据')
    env_raw = pd.read_excel(env_path, sheet_name=cfg['env_sheet'])
    env = pd.DataFrame({'timestamp': pd.to_datetime(env_raw[cfg['env_time_column']], errors='raise')})
    for short, col in cfg['environment_columns'].items():
        env[short] = pd.to_numeric(env_raw[col], errors='raise')
    if env.timestamp.isna().any() or env.timestamp.duplicated().any():
        raise ValueError('环境时间缺失或重复，请先检查原始数据')
    return data.sort_values('timestamp'), env.sort_values('timestamp'), mapping


def write_sheets(path, sheets):
    """Create result workbooks or append sheets to copied originals."""
    options = {'mode': 'a', 'if_sheet_exists': 'replace'} if Path(path).exists() else {'mode': 'w'}
    with pd.ExcelWriter(path, engine='openpyxl', **options) as writer:
        for name, frame in sheets.items():
            frame.to_excel(writer, sheet_name=name, index=False)
    wb = load_workbook(path)
    for name in sheets:
        ws = wb[name]
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        ws.row_dimensions[1].height = 30
        for c in ws[1]:
            c.font = Font(color='FFFFFF', bold=True)
            c.fill = PatternFill('solid', fgColor='24475B')
            c.alignment = Alignment(wrap_text=True, vertical='center')
        for col in ws.columns:
            values = [str(c.value or '') for c in list(col)[:80]]
            display_width = lambda s: sum(2 if unicodedata.east_asian_width(ch) in 'WF' else 1 for ch in s)
            width = min(65, max(15, max(map(display_width, values)) + 3))
            ws.column_dimensions[col[0].column_letter].width = width
            for cell in list(col)[1:]:
                if cell.is_date:
                    cell.number_format = 'yyyy-mm-dd hh:mm:ss'
                elif isinstance(cell.value, float):
                    if col[0].value in ['p','q_BH']:
                        cell.number_format = '0.000E+00'
                    elif col[0].value == 'factor_nmol_m2_s_per_ppb':
                        cell.number_format = '0.000000000000'
                    else:
                        cell.number_format = '0.000000E+00' if 0 < abs(cell.value) < .0001 else '0.000000'
                elif isinstance(cell.value, str):
                    cell.alignment = Alignment(wrap_text=True, vertical='top')
                    height = 16 * (1 + display_width(cell.value)//max(1,int(width)-2))
                    ws.row_dimensions[cell.row].height = max(ws.row_dimensions[cell.row].height or 16, height)
    wb.save(path)


def plot_lines(frame, variables, target, unit, title):
    """Keep the separate category/environment figures alongside the combined figure."""
    fig, axes = plt.subplots(len(variables), 1, figsize=(10, 2.5*len(variables)), squeeze=False)
    for ax, variable in zip(axes[:, 0], variables):
        group = frame[frame.variable.eq(variable)].sort_values('hour')
        ax.errorbar(group.hour, group['mean'], yerr=group.se, fmt='o-', capsize=3, linewidth=1.4)
        ax.set_ylabel(PLOT_LABELS.get(variable, variable) + '\n' + unit)
        ax.xaxis.set_major_formatter(mdates.DateFormatter('%H:%M'))
        ax.grid(alpha=.22)
    axes[0, 0].set_title(title + ' (mean ± SE)')
    axes[-1, 0].set_xlabel('Sampling hour (UTC+8)')
    fig.tight_layout()
    fig.savefig(target.with_suffix('.png'), dpi=180)
    plt.close(fig)


def plot_combined(categories, environment, target, unit='ppb', quantity='VOC concentration'):
    """Nine panels with a common time axis and independent, labeled units."""
    panels = [(environment, 'temperature', 'Temperature', '°C', '#D65A3A'),
              (environment, 'humidity', 'Relative humidity', '%', '#3B7CA8'),
              (environment, 'PAR', 'PAR', 'μmol/m²/s', '#B58B16')]
    titles = [PLOT_LABELS[category] for category in CATEGORIES]
    colors = ['#287D8E', '#3D8C57', '#8D6BA4', '#4269A0', '#AA6A43', '#B24F77']
    panels += [(categories, category, title, unit, color)
               for category, title, color in zip(CATEGORIES, titles, colors)]
    fig, axes = plt.subplots(3, 3, figsize=(16, 11.5), sharex=True, layout='constrained')
    all_hours = pd.to_datetime(categories.hour)
    date_label = all_hours.min().strftime('%Y-%m-%d')
    multiple_days = all_hours.dt.normalize().nunique() > 1
    if multiple_days:
        date_label += ' to ' + all_hours.max().strftime('%Y-%m-%d')
    for ax, (frame, variable, title, axis_unit, color) in zip(axes.flat, panels):
        g = frame[frame.variable.eq(variable)].sort_values('hour')
        ax.errorbar(g.hour, g['mean'], yerr=g.se, fmt='o-', capsize=3,
                    linewidth=1.6, markersize=4, color=color)
        ax.set_title(title, fontsize=12, pad=10)
        ax.set_ylabel(axis_unit)
        ax.xaxis.set_major_locator(mdates.HourLocator(interval=2))
        ax.xaxis.set_major_formatter(mdates.DateFormatter('%m-%d\n%H:%M' if multiple_days else '%H:%M'))
        ax.tick_params(axis='x', labelbottom=True)
        ax.tick_params(labelsize=9)
        ax.ticklabel_format(axis='y', style='sci', scilimits=(-3, 4), useOffset=False)
        ax.grid(alpha=.22)
    for ax in axes[-1]:
        ax.set_xlabel('Sampling hour (UTC+8)')
    fig.suptitle(f'{date_label} {quantity} and environmental parameters (mean ± SE)', fontsize=16)
    fig.savefig(target.with_suffix('.png'), dpi=180)
    plt.close(fig)


def plot_correlation(correlations_table, target, method, quantity='concentration'):
    matrix = correlations_table[correlations_table.method.eq(method)].pivot(
        index='variable', columns='environment', values='r').reindex(CATEGORIES)
    fig, ax = plt.subplots(figsize=(6, 5))
    image = ax.imshow(np.ma.masked_invalid(matrix.values), vmin=-1, vmax=1, cmap='RdBu_r')
    labels = [PLOT_LABELS.get(c, c) for c in matrix.columns]
    ax.set_xticks(range(len(matrix.columns)), labels, fontsize=9, rotation=30,
                  ha='right', rotation_mode='anchor')
    ax.set_yticks(range(6), CATEGORIES)
    ax.set_xlabel('Environmental parameter')
    ax.set_ylabel('VOC category')
    for i in range(6):
        for j in range(len(matrix.columns)):
            value = matrix.iloc[i,j]
            ax.text(j,i,f'{value:.2f}' if pd.notna(value) else 'NA',ha='center',va='center')
    ax.set_title(f'{method}: {quantity} vs environment')
    fig.colorbar(image, ax=ax, label='Correlation coefficient')
    fig.tight_layout()
    fig.savefig(target.with_suffix('.png'), dpi=180)
    plt.close(fig)


def run(voc_path, env_path, output, config=None, mapping_path=None):
    cfg = config if isinstance(config, dict) else json.loads(Path(config or ROOT/'config.json').read_text(encoding='utf-8'))
    if cfg['window_minutes'] <= 0:
        raise ValueError('window_minutes 必须大于0')
    out = Path(output)
    if out.exists() and any(out.iterdir()):
        raise ValueError('输出目录非空，请使用新的目录，防止覆盖既有结果')
    data, env, mapping = read_inputs(voc_path, env_path, cfg, mapping_path or ROOT/'species_mapping.csv')
    # Only the six requested classes enter any derived analysis output.
    mapping = mapping.loc[mapping.category.isin(CATEGORIES)].copy()
    if mapping.empty:
        raise ValueError('没有可归入六类的物种，请检查分类表')
    species = mapping.species.tolist()
    hourly = summarize(data, ['hour'], species)
    totals = category_totals(data, mapping)
    cat_hourly = summarize(totals, ['hour'], CATEGORIES)
    windows, matched = match_environment(data, env, cfg['window_mode'], cfg['window_minutes'])
    env_stats = summarize(matched, ['hour'], list(cfg['environment_columns']))
    # Reindex to retain hours with no matched environmental data.
    idx = pd.MultiIndex.from_product([data.hour.unique(), list(cfg['environment_columns'])], names=['hour','variable'])
    env_stats = env_stats.set_index(['hour','variable']).reindex(idx).reset_index()
    env_stats['n'] = env_stats['n'].fillna(0).astype(int)
    settings = flux_settings(voc_path, cfg)
    flux_rows, flux_parameters = calculate_flux_rows(data, species, env_stats, settings)
    flux_hourly = summarize(flux_rows, ['hour'], species)
    flux_totals = category_totals(flux_rows, mapping)
    flux_cat_hourly = summarize(flux_totals, ['hour'], CATEGORIES)
    for frame in [flux_hourly, flux_cat_hourly]:
        frame['unit'] = 'nmol m-2 s-1'
    corr_cat = correlations(flux_cat_hourly, env_stats, 'category')
    corr_species = correlations(flux_hourly, env_stats, 'species')
    corr_cat['basis'] = 'leaf_area_flux_nmol_m2_s'
    corr_species['basis'] = 'leaf_area_flux_nmol_m2_s'
    metadata = data.groupby('hour').timestamp.agg(sample_start='min', sample_end='max', sample_n='count').reset_index()
    aligned = cat_hourly.pivot(index='hour', columns='variable', values='mean').join(
        env_stats.pivot(index='hour', columns='variable', values='mean')).reset_index()
    aligned_flux = flux_cat_hourly.pivot(index='hour', columns='variable', values='mean').join(
        env_stats.pivot(index='hour', columns='variable', values='mean')).reset_index()
    notes = [
        ('SE', '样本标准差(ddof=1)/sqrt(n)；n<2时SE为空；n为技术读数数，不代表生物学重复。'),
        ('类别总量', '每条记录先合计同类物种，再按小时计算mean/SE；成员缺失时该条总量缺失；无成员类别留空。'),
        ('时间', 'hour为采样所在整点分组，非整小时连续测量；保留sample_start/end。时区沿用China Standard Time。'),
        ('环境窗口', f"{cfg['window_mode']}, {cfg['window_minutes']}分钟；每小时取各采样窗口并集，环境行去重。五分钟区间左闭右开；±窗口两端包含。"),
        ('相关性', '对齐小时mean，成对删除缺失；Pearson/Spearman双侧p；BH按层级和方法分别对全部有效检验校正。'),
        ('局限', f'仅{data.hour.nunique()}个时间点的探索分析；时间自相关、日变化及技术重复影响推断，相关性不代表因果。'),
        ('浓度单位', cfg['concentration_unit']),
        ('通量单位', 'nmol m-2 s-1；正值为净释放，负值为净吸收；不得把所有成分都直接认定为生物源。'),
        ('通量公式', 'F=(Cout-Cin)[ppb] × Q[L/min] × 0.001 × Pref[Pa] / (60 × R × Tref[K] × A[m²])；R=8.31446261815324。'),
        ('实验条件', f"叶室{settings['chamber_volume_l']} L；总流量{settings['flow_l_min']} L/min；叶面积{settings['leaf_area_m2']} m²；稳态已由实验者确认。"),
        ('流量基准', f"{settings['flow_basis']}；参考温度{settings.get('reference_temperature_c')} °C；参考压力{settings.get('reference_pressure_pa')} Pa。标准流量不再按环境温度修正。"),
        ('背景', f"{settings['background_mode']}：已扣背景的数据不再次扣除；负通量保留。"),
        ('体积作用', '稳态F不直接乘叶室体积；V/Q约2 min为参考基准下的名义时间常数，不是稳态证据；未做储存项、壁面损失、化学反应或水汽稀释修正。'),
        ('通量SE', '先逐记录换算通量，再按物种或类别汇总；SE仅反映技术读数离散，不含流量、温压、背景及面积测量不确定度。'),
        ('相关性对象', '当前类别和物种相关性使用单位叶面积通量与环境参数的小时均值。'),
        ('方法来源', 'https://bg.copernicus.org/articles/6/405/2009/；https://amt.copernicus.org/articles/14/4445/2021/'),
        ('范围', '仅分析已归入iso、MTs、SQTs、OTs、OMTs、OVOCs的物种；六类之外的成分直接剔除，不进入新统计sheet、绘图或相关性。'),
        ('OMTs依据', '芳樟醇：https://www.kegg.jp/entry/C11388；分类表是可编辑的分析口径。')]
    out.mkdir(parents=True, exist_ok=True)
    voc_out, env_out = out/(Path(voc_path).stem+'_analysis.xlsx'), out/(Path(env_path).stem+'_analysis.xlsx')
    for src, dst in [(voc_path, voc_out), (env_path, env_out)]:
        if Path(src).resolve() == dst.resolve():
            raise ValueError('不可覆盖原始文件')
        shutil.copy2(src, dst)
    write_sheets(voc_out, {'小时通量统计': flux_hourly, '六类通量统计': flux_cat_hourly,
        '通量逐记录明细': flux_rows, '六类逐记录通量': flux_totals,
        '通量换算参数': flux_parameters, '对齐小时通量': aligned_flux,
        '小时浓度统计': hourly, '六类小时统计': cat_hourly,
        '分类逐记录合计': totals, '物种分类表': mapping, '采样时段': metadata,
        '对齐小时均值': aligned, '类别相关性': corr_cat, '物种相关性': corr_species,
        '方法说明': pd.DataFrame(notes, columns=['item','description'])})
    write_sheets(env_out, {'匹配环境统计': env_stats, '匹配环境原始行': matched,
        '匹配窗口': windows, '采样时段': metadata,
        '方法说明': pd.DataFrame(notes, columns=['item','description'])})
    for name, frame in [('hourly_species', hourly), ('hourly_categories', cat_hourly),
                        ('hourly_environment', env_stats), ('correlation_categories', corr_cat),
                        ('correlation_species', corr_species), ('aligned_means', aligned),
                        ('hourly_flux_species', flux_hourly), ('hourly_flux_categories', flux_cat_hourly),
                        ('flux_conversion_parameters', flux_parameters), ('aligned_flux_means', aligned_flux)]:
        frame.to_csv(out/(name+'.csv'), index=False, encoding='utf-8-sig')
    plots = out/'plots'
    plots.mkdir(exist_ok=True)
    plt.rcParams['font.sans-serif'] = ['DejaVu Sans']
    plt.rcParams['axes.unicode_minus'] = False
    plot_combined(flux_cat_hourly, env_stats, plots/'categories_environment', FLUX_UNIT, 'Leaf-area BVOC flux')
    plot_lines(flux_cat_hourly, CATEGORIES, plots/'categories', FLUX_UNIT, 'Leaf-area BVOC flux')
    for variable, unit in [('temperature','°C'),('humidity','%'),('PAR','μmol/m²/s')]:
        plot_lines(env_stats, [variable], plots/variable, unit, 'Matched environment')
    for method in ['Pearson','Spearman']:
        plot_correlation(corr_cat, plots/f'correlation_{method}', method, 'leaf-area flux')
    manifest = {'config': cfg, 'inputs': {str(Path(p).name): hashlib.sha256(Path(p).read_bytes()).hexdigest()
                for p in [voc_path, env_path]}, 'hour_count': int(data.hour.nunique()),
                'sample_count': len(data), 'category_members': mapping.category.value_counts().to_dict(),
                'analysis_kind': 'steady_state_leaf_area_flux', 'dataset_folder': Path(voc_path).parent.name,
                'flux_settings': settings, 'flux_unit': 'nmol m-2 s-1',
                'mapping_sha256': hashlib.sha256(Path(mapping_path or ROOT/'species_mapping.csv').read_bytes()).hexdigest(),
                'unmatched_hours': env_stats.loc[env_stats.n.eq(0), 'hour'].astype(str).unique().tolist(),
                'versions': {m: __import__(m).__version__ for m in ['pandas','numpy','scipy','matplotlib','openpyxl']}}
    (out/'manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf-8')
    mapping.to_csv(out/'species_mapping.csv', index=False, encoding='utf-8-sig')
    return manifest


def main(argv=None):
    parser = argparse.ArgumentParser(description='VOCs小时统计、环境匹配、绘图和相关性分析')
    source = parser.add_mutually_exclusive_group()
    source.add_argument('--dataset', help='Original Data 下的数据文件夹名，例如 1015_xiangzhang')
    source.add_argument('--input-dir', type=Path, help='直接指定包含两个源Excel的文件夹')
    source.add_argument('--all', action='store_true', help='批量分析所有原始数据子文件夹，跳过已有结果')
    parser.add_argument('--input-root', type=Path, default=INPUT_ROOT)
    parser.add_argument('--voc', type=Path)
    parser.add_argument('--env', type=Path)
    parser.add_argument('--output', type=Path, help='默认 results/<源浓度Excel文件名（无扩展名）>')
    parser.add_argument('--config', type=Path, default=ROOT/'config.json')
    parser.add_argument('--mapping', type=Path, default=ROOT/'species_mapping.csv')
    parser.add_argument('--leaf-area', type=float, help='叶面积m²；通常从数据集文件夹末尾_0.5自动识别')
    parser.add_argument('--check-inputs', action='store_true', help='只验证文件识别和数据读取，不创建结果')
    args = parser.parse_args(argv)
    explicit = args.voc is not None or args.env is not None
    if explicit:
        if args.voc is None or args.env is None or args.dataset or args.input_dir or args.all:
            parser.error('--voc 和 --env 须一起使用，不能与 --dataset/--input-dir/--all 混用')
        jobs = [(args.voc, args.env)]
    else:
        if args.input_dir:
            folders = [args.input_dir]
        elif args.dataset:
            if Path(args.dataset).name != args.dataset:
                parser.error('--dataset 只填子文件夹名；完整路径使用 --input-dir')
            folders = [args.input_root / args.dataset]
        else:
            folders = sorted(p for p in args.input_root.iterdir() if p.is_dir()) if args.input_root.is_dir() else []
            if not folders:
                parser.error(f'没有原始数据文件夹: {args.input_root}')
            if not args.all and len(folders) != 1:
                parser.error('发现多个日期，请用 --dataset 指定一个文件夹，或用 --all 批量处理')
        jobs = [find_source_pair(folder) for folder in folders]
    if args.all and args.output:
        parser.error('批量模式使用各源文件同名结果目录，不能指定统一的 --output')
    cfg = json.loads(args.config.read_text(encoding='utf-8'))
    if args.leaf_area is not None:
        cfg.setdefault('flux', {})['leaf_area_m2'] = args.leaf_area
    for voc_path, env_path in jobs:
        output = args.output or ROOT/'results'/voc_path.stem
        if args.check_inputs:
            data, environment, _ = read_inputs(voc_path, env_path, cfg, args.mapping)
            settings = flux_settings(voc_path, cfg)
            print(json.dumps({'voc': str(voc_path.resolve()), 'env': str(env_path.resolve()),
                              'output': str(output.resolve()), 'samples': len(data),
                              'hours': int(data.hour.nunique()), 'environment_records': len(environment),
                              'flux_settings': settings},
                             ensure_ascii=False, indent=2))
        elif args.all and output.exists() and any(output.iterdir()):
            manifest_path = output/'manifest.json'
            existing = json.loads(manifest_path.read_text(encoding='utf-8')) if manifest_path.exists() else {}
            expected_inputs = {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in [voc_path, env_path]}
            matches = (existing.get('analysis_kind') == 'steady_state_leaf_area_flux'
                       and existing.get('flux_settings') == flux_settings(voc_path, cfg)
                       and existing.get('inputs') == expected_inputs and existing.get('config') == cfg
                       and existing.get('mapping_sha256') == hashlib.sha256(args.mapping.read_bytes()).hexdigest())
            if not matches:
                raise ValueError(f'已有结果与当前源文件/叶面积/分析参数不一致: {output}；'
                                 '请用 --dataset 指定该组，并用 --output 指定新的结果子目录重新分析')
            print(f'跳过已完成且源文件/参数一致的通量结果: {output}')
        else:
            print(json.dumps(run(voc_path, env_path, output, cfg, args.mapping), ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
